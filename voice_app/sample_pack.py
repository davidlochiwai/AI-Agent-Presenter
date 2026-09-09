"""Write sample_talk.pptx, script.xlsx, and knowledge_sources/knowledge.xlsx from talk_content."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from voice_app.talk_content import KNOWLEDGE, SCRIPT, SLIDES

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
DECK_PATH = DATA / "sample_talk.pptx"
SCRIPT_PATH = DATA / "script.xlsx"
KNOWLEDGE_DIR = DATA / "knowledge_sources"
KNOWLEDGE_PATH = KNOWLEDGE_DIR / "knowledge.xlsx"

NAVY = RGBColor(0x1B, 0x3A, 0x4B)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
INK = RGBColor(0x1A, 0x24, 0x2A)
MUTED = RGBColor(0x5A, 0x6A, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF7, 0xF4, 0xEE)
FONT = "Microsoft JhengHei"


def _set_run(paragraph, text: str, size: int, color: RGBColor, bold: bool = False) -> None:
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = FONT


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _textbox(slide, left, top, width, height, text: str, size: int, color: RGBColor, bold: bool = False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p, text, size, color, bold)
    return box


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _footer(slide, index: int, total: int) -> None:
    _textbox(
        slide,
        Inches(0.7),
        Inches(7.05),
        Inches(10.5),
        Inches(0.3),
        "Harbour AI  ·  2026 Q2 內部匯報  ·  資料截至 6 月 30 日",
        11,
        MUTED,
    )
    _textbox(
        slide,
        Inches(11.4),
        Inches(7.05),
        Inches(1.2),
        Inches(0.3),
        f"{index} / {total}",
        11,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )


def _add_title_slide(prs: Presentation, spec: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    _fill(bg, NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(2.35), Inches(2.2), Inches(0.08))
    _fill(accent, GOLD)
    _textbox(slide, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5), spec.get("kicker") or "", 16, GOLD, bold=True)
    _textbox(slide, Inches(0.7), Inches(2.55), Inches(12), Inches(1.4), spec["title"], 54, WHITE, bold=True)
    _textbox(slide, Inches(0.7), Inches(4.2), Inches(12), Inches(1.0), spec.get("subtitle") or "", 24, CREAM)
    _notes(slide, spec.get("notes") or "")
    _textbox(
        slide,
        Inches(0.7),
        Inches(7.05),
        Inches(10.5),
        Inches(0.3),
        "Harbour AI  ·  2026 Q2 內部匯報",
        11,
        RGBColor(0xA8, 0xB8, 0xC0),
    )
    _textbox(
        slide,
        Inches(11.4),
        Inches(7.05),
        Inches(1.2),
        Inches(0.3),
        f"{index} / {total}",
        11,
        RGBColor(0xA8, 0xB8, 0xC0),
        align=PP_ALIGN.RIGHT,
    )


def _add_content_slide(prs: Presentation, spec: dict, index: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    _fill(bg, CREAM)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), prs.slide_width, Inches(1.35))
    _fill(bar, NAVY)
    gold = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Inches(1.35), prs.slide_width, Inches(0.06))
    _fill(gold, GOLD)
    _textbox(slide, Inches(0.7), Inches(0.35), Inches(12), Inches(0.8), spec["title"], 28, WHITE, bold=True)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(12), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    bullets = spec.get("bullets") or []
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(14)
        _set_run(p, "•  " + line, 22, INK, bold=False)
    _notes(slide, spec.get("notes") or "")
    _footer(slide, index, total)


def write_sample_deck(path: Path | None = None) -> Path:
    path = Path(path or DECK_PATH)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = len(SLIDES)
    for i, spec in enumerate(SLIDES, start=1):
        if spec.get("layout") == "title":
            _add_title_slide(prs, spec, i, total)
        else:
            _add_content_slide(prs, spec, i, total)
    prs.save(str(path))
    return path


def _style_header(ws, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="1B3A4B")
    font = Font(name="Calibri", bold=True, color="FFFFFF")
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(1, col, name)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _write_sheet(ws, rows: list[dict], headers: list[str], widths: dict[str, int]) -> None:
    wrap = Alignment(wrap_text=True, vertical="top")
    body = Font(name="Calibri", size=11)
    _style_header(ws, headers)
    for r, row in enumerate(rows, start=2):
        for c, key in enumerate(headers, start=1):
            value = row.get(key, "")
            if value is None:
                value = ""
            cell = ws.cell(r, c, value)
            cell.font = body
            cell.alignment = wrap
        ws.row_dimensions[r].height = 48
    for c, key in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(key, 18)
    ws.row_dimensions[1].height = 22


def write_script_xlsx(path: Path | None = None) -> Path:
    path = Path(path or SCRIPT_PATH)
    wb = Workbook()
    ws = wb.active
    ws.title = "script"
    headers = ["seq", "slide", "actions", "script_yue", "script_en", "notes"]
    _write_sheet(ws, SCRIPT, headers, {
        "seq": 8, "slide": 8, "actions": 16, "script_yue": 55, "script_en": 55, "notes": 28,
    })
    note = wb.create_sheet("readme")
    note["A1"] = "The agent reads every beat in order unless the audience interrupts. There is no ask-to-continue pause."
    note["A2"] = "seq and slide must match sample_talk.pptx. Regenerated by: python -m voice_app.sample_pack"
    note.column_dimensions["A"].width = 120
    wb.save(path)
    return path


def write_knowledge_xlsx(path: Path | None = None) -> Path:
    path = Path(path or KNOWLEDGE_PATH)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "knowledge"
    headers = ["card_id", "slide", "topic", "keywords", "fact_yue", "fact_en"]
    rows = [dict(card) for card in KNOWLEDGE]
    _write_sheet(ws, rows, headers, {
        "card_id": 16, "slide": 10, "topic": 22, "keywords": 40, "fact_yue": 55, "fact_en": 55,
    })
    note = wb.create_sheet("readme")
    note["A1"] = "Empty slide = answer without flipping PowerPoint. A number = go to that 1-based slide when this card is used."
    note["A2"] = "Q&A uses these approved facts as retrieval evidence. Unsupported or failed retrieval returns a safe abstention."
    note.column_dimensions["A"].width = 120
    wb.save(path)
    return path


def write_all() -> dict[str, Path]:
    DATA.mkdir(parents=True, exist_ok=True)
    if len(SLIDES) != len(SCRIPT):
        raise SystemExit(f"Slide count {len(SLIDES)} != script beats {len(SCRIPT)}")
    for row in SCRIPT:
        if not 1 <= int(row["slide"]) <= len(SLIDES):
            raise SystemExit(f"Script seq {row['seq']} points at missing slide {row['slide']}")
    written: dict[str, Path] = {}
    for key, fn in (("deck", write_sample_deck), ("script", write_script_xlsx), ("knowledge", write_knowledge_xlsx)):
        try:
            written[key] = fn()
        except PermissionError:
            print(f"Skip {key}: file is open in another program. Close it and re-run python -m voice_app.sample_pack")
    for stale in (DATA / "script.csv", DATA / "knowledge.csv", DATA / "knowledge.xlsx"):
        if stale.exists():
            stale.unlink()
            written[stale.name] = stale
    return written


def main() -> None:
    for label, path in write_all().items():
        print(f"{label}: {path}")
    print(f"{len(SLIDES)} slides, {len(SCRIPT)} beats, {len(KNOWLEDGE)} knowledge cards.")


if __name__ == "__main__":
    main()
