"""Load presenter beats from data/script.xlsx, with the bundled sample as fallback."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from voice_app import config
from voice_app.talk_content import SCRIPT


def _clean(value: Any) -> str:
    return str(value or "").strip()


def load_script_beats(path: Path | None = None) -> tuple[list[dict[str, Any]], str]:
    """Return (beats, source_label). Excel wins when the file exists and has rows."""
    workbook_path = Path(path or config.SAMPLE_SCRIPT_PATH)
    if not workbook_path.exists():
        return [dict(row) for row in SCRIPT], "bundled sample (talk_content.py)"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["script"] if "script" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [_clean(value) for value in next(rows, [])]
        beats: list[dict[str, Any]] = []
        for values in rows:
            row = dict(zip(headers, values))
            spoken = _clean(row.get("script_yue"))
            if not spoken:
                continue
            try:
                seq = int(row.get("seq") or 0)
            except (TypeError, ValueError):
                seq = len(beats) + 1
            try:
                slide = int(row.get("slide") or 0)
            except (TypeError, ValueError):
                slide = 0
            if slide < 1:
                continue
            actions = _clean(row.get("actions")) or f"goto_slide:{slide}"
            beats.append(
                {
                    "seq": seq if seq > 0 else len(beats) + 1,
                    "slide": slide,
                    "actions": actions,
                    "script_yue": spoken,
                    "script_en": _clean(row.get("script_en")),
                    "notes": _clean(row.get("notes")),
                }
            )
    finally:
        workbook.close()
    beats.sort(key=lambda item: (int(item["seq"]), int(item["slide"])))
    if not beats:
        return [dict(row) for row in SCRIPT], "bundled sample (empty script.xlsx)"
    return beats, str(workbook_path)
